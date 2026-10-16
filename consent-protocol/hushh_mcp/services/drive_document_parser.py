"""Bounded, credential-free parsing leaf; executed in a disposable process.

Never follows links, renders scripts, extracts archives to disk, or performs OCR.
Input is private and must not be logged, even on malformed parser exceptions.
"""

from __future__ import annotations

import csv
import io
import json
import re
import sys
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePosixPath

from defusedxml import ElementTree

try:
    import resource
except ImportError:  # Windows has no POSIX resource-limit module.
    resource = None

MAX_INPUT_BYTES = 4 * 1024 * 1024
MAX_TEXT_BYTES = 256 * 1024
MAX_PAGES = 100
# Live lane only: a Google Sheets CSV export or an uploaded .csv.
MAX_CSV_ROWS = 2000
MAX_CSV_COLUMNS = 64
MAX_CSV_CELL_CHARS = 1000
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_XLSX_INPUT_BYTES = 10 * 1024 * 1024
MAX_XLSX_EXPANDED_BYTES = 32 * 1024 * 1024
MAX_XLSX_SHEETS = 64
MAX_XLSX_ROWS = 20_000
MAX_XLSX_COLUMNS = 256
MAX_XLSX_CELLS = 150_000


class ParseError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedText:
    pages: tuple[str, ...] = field(repr=False)
    truncated: bool = False


def _bounded(pages) -> ParsedText:
    kept: list[str] = []
    remaining = MAX_TEXT_BYTES
    truncated = False
    for number, value in enumerate(pages):
        if number >= MAX_PAGES:
            truncated = True
            break
        value = value.replace("\x00", "")
        encoded = value.encode("utf-8")
        if len(encoded) > remaining:
            kept.append(encoded[:remaining].decode("utf-8", errors="ignore"))
            truncated = True
            break
        kept.append(value)
        remaining -= len(encoded)
    if not any(page.strip() for page in kept):
        raise ParseError("no_extractable_text")
    return ParsedText(tuple(kept), truncated)


def _docx(content: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(item.file_size for item in entries) > 16 * 1024 * 1024:
            raise ParseError("file_too_large")
        names = [item.filename for item in entries]
        if len(names) != len(set(names)):
            raise ParseError("invalid_document")
        for item in entries:
            path = PurePosixPath(item.filename)
            if (
                item.flag_bits & 1
                or path.is_absolute()
                or ".." in path.parts
                or "\\" in item.filename
                or item.file_size > max(1024 * 1024, item.compress_size * 100)
                or "vbaproject" in item.filename.lower()
                or item.filename.startswith("word/embeddings/")
            ):
                raise ParseError("invalid_document")
        if "[Content_Types].xml" not in names or "word/document.xml" not in names:
            raise ParseError("invalid_document")
        # No relationship is dereferenced, including external image/template URLs.
        root = ElementTree.fromstring(archive.read("word/document.xml"))
        namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        lines = []
        for paragraph in root.iter(namespace + "p"):
            lines.append("".join(node.text or "" for node in paragraph.iter(namespace + "t")))
        return "\n".join(lines)


def parse_document(content: bytes, mime_type: str) -> ParsedText:
    if not content or len(content) > MAX_INPUT_BYTES:
        raise ParseError("file_too_large" if content else "no_extractable_text")
    try:
        if mime_type in {"text/plain", "text/markdown"}:
            return _bounded((content.decode("utf-8-sig", errors="strict"),))
        if mime_type == DOCX:
            return _bounded((_docx(content),))
        if mime_type == "application/pdf":
            import pdfplumber
            from pdfminer.pdfpage import PDFPage
            from pdfplumber.page import Page

            if not content.startswith(b"%PDF-"):
                raise ParseError("invalid_document")
            pdf = pdfplumber.open(io.BytesIO(content))
            try:
                if pdf.doc.encryption:
                    raise ParseError("encrypted_document")

                def page_text():
                    for number, source_page in enumerate(PDFPage.create_pages(pdf.doc), 1):
                        if number > MAX_PAGES:
                            yield ""  # Overflow marker; never extract page 101.
                            break
                        page = Page(pdf, source_page, page_number=number)
                        try:
                            value = page.extract_text() or ""
                            # Compatibility readers do not run OCR. A mixed
                            # PDF must not disguise image-only evidence as an
                            # empty, successfully extracted page.
                            if page.images or (not value.strip() and page.objects):
                                missing.append(number)
                            yield value
                        finally:
                            page.close()

                missing: list[int] = []
                parsed = _bounded(page_text())
                return ParsedText(parsed.pages, parsed.truncated or bool(missing))
            finally:
                # PDF.close()/__exit__ accesses .pages and materializes the
                # entire page tree. Close the stream and caches directly.
                pdf.flush_cache()
                pdf.stream.close()
        raise ParseError("unsupported_format")
    except ParseError:
        raise
    except Exception:
        raise ParseError("invalid_document") from None


def parse_live_csv(content: bytes) -> ParsedText:
    """Bounded CSV rows as plain ``a | b`` text lines for the live lane.

    Cells are data, never formulas or links: nothing is evaluated. Stops at the
    row, column, cell and text caps and fails closed on anything not clean
    UTF-8 CSV. Never calls csv.field_size_limit, which is process-global; a
    field over its default limit is a damaged file here.
    """
    if not content or len(content) > MAX_INPUT_BYTES:
        raise ParseError("file_too_large" if content else "no_extractable_text")
    try:
        text = content.decode("utf-8-sig", errors="strict")
    except UnicodeDecodeError:
        raise ParseError("invalid_document") from None
    if "\x00" in text:
        raise ParseError("invalid_document")
    lines: list[str] = []
    size = 0
    truncated = False
    try:
        for number, row in enumerate(csv.reader(io.StringIO(text, newline=""), strict=True)):
            if number >= MAX_CSV_ROWS or size > MAX_TEXT_BYTES:
                truncated = True
                break
            cells = [" ".join(cell.split()) for cell in row]
            if any(cells[MAX_CSV_COLUMNS:]):
                truncated = True
            cells = cells[:MAX_CSV_COLUMNS]
            while cells and not cells[-1]:
                cells.pop()  # Sheets pads short rows with empty cells.
            if not cells:
                continue
            if any(len(cell) > MAX_CSV_CELL_CHARS for cell in cells):
                truncated = True
                cells = [cell[:MAX_CSV_CELL_CHARS] for cell in cells]
            line = " | ".join(cells)
            lines.append(line)
            size += len(line.encode("utf-8")) + 1
    except csv.Error:
        raise ParseError("invalid_document") from None
    parsed = _bounded(("\n".join(lines),))
    return ParsedText(parsed.pages, parsed.truncated or truncated)


def parse_live_xlsx(content: bytes) -> ParsedText:
    """Read every worksheet from bounded original bytes, including hidden tabs.

    Cached formula results can be used as observations, but a formula without
    a cached result or any budget stop marks the extraction incomplete.
    """
    if not content or len(content) > MAX_XLSX_INPUT_BYTES:
        raise ParseError("file_too_large" if content else "no_extractable_text")
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter

        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            names = [item.filename for item in entries]
            if (
                len(entries) > 2000
                or len(names) != len(set(names))
                or sum(item.file_size for item in entries) > MAX_XLSX_EXPANDED_BYTES
                or "[Content_Types].xml" not in names
                or "xl/workbook.xml" not in names
            ):
                raise ParseError("invalid_document")
            for item in entries:
                path = PurePosixPath(item.filename)
                if (
                    path.is_absolute()
                    or ".." in path.parts
                    or "\\" in item.filename
                    or item.flag_bits & 1
                    or item.file_size > MAX_XLSX_EXPANDED_BYTES
                    or item.file_size > max(1024 * 1024, item.compress_size * 100)
                    or "vbaproject" in item.filename.lower()
                ):
                    raise ParseError("invalid_document")
            # Formula cells with no cached value are absent from data_only
            # readers. Inspect the bounded sheet XML so that absence cannot be
            # mistaken for a complete empty worksheet.
            missing_formula_values = False
            namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
            for item in entries:
                if not (
                    item.filename.startswith("xl/worksheets/sheet")
                    and item.filename.endswith(".xml")
                ):
                    continue
                root = ElementTree.fromstring(archive.read(item))
                for cell in root.iter(namespace + "c"):
                    children = list(cell)
                    if any(child.tag == namespace + "f" for child in children) and not any(
                        child.tag == namespace + "v" and (child.text or "").strip()
                        for child in children
                    ):
                        missing_formula_values = True
                        break
                if missing_formula_values:
                    break
            partial = missing_formula_values or any(
                name.startswith(("xl/externalLinks/", "xl/drawings/", "xl/charts/"))
                for name in names
            )
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True, keep_links=False
        )
        try:
            sheets = workbook.worksheets
            partial = partial or len(sheets) > MAX_XLSX_SHEETS
            lines: list[str] = []
            used_bytes = 0
            populated = False
            cells_seen = 0
            budget_stop = False
            for sheet in sheets[:MAX_XLSX_SHEETS]:
                sheet_populated = False
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    if row_number > MAX_XLSX_ROWS:
                        partial = budget_stop = True
                        break
                    if len(row) > MAX_XLSX_COLUMNS:
                        partial = True
                    cells = []
                    for column_number, value in enumerate(row[:MAX_XLSX_COLUMNS], 1):
                        cells_seen += 1
                        if cells_seen > MAX_XLSX_CELLS:
                            partial = budget_stop = True
                            break
                        if value is None:
                            continue
                        rendered = value.isoformat() if hasattr(value, "isoformat") else str(value)
                        rendered = " ".join(rendered.split())
                        if rendered:
                            if len(rendered) > MAX_CSV_CELL_CHARS:
                                rendered = rendered[:MAX_CSV_CELL_CHARS]
                                partial = True
                            cells.append(f"{get_column_letter(column_number)}={rendered}")
                    if cells_seen > MAX_XLSX_CELLS:
                        break
                    if not cells:
                        continue
                    populated = sheet_populated = True
                    line = f"[{sheet.title}] row {row_number}: " + " | ".join(cells)
                    line_bytes = len(line.encode("utf-8")) + 1
                    if used_bytes + line_bytes > MAX_TEXT_BYTES:
                        partial = budget_stop = True
                        break
                    lines.append(line)
                    used_bytes += line_bytes
                if not sheet_populated and not budget_stop:
                    lines.append(f"[{sheet.title}] empty")
                if budget_stop:
                    break
            if not populated:
                raise ParseError("no_extractable_text")
            parsed = _bounded(("\n".join(lines),))
            return ParsedText(parsed.pages, parsed.truncated or partial)
        finally:
            workbook.close()
    except ParseError:
        raise
    except Exception:
        raise ParseError("invalid_document") from None


# A Google Docs Markdown export inlines every image as a base64 data URI, as a
# reference definition ("[image1]: <data:image/png;base64,...>") used by
# "![][image1]", or inline. Those bytes are not text: they would spend the text
# budget and reach the model as noise. Each image becomes one "[image]" marker.
_DATA_URI_DEFINITION = re.compile(
    rb"^[ \t]*\[[^\]\n]{1,200}\]:[ \t]*<?data:[^\s>]*>?[ \t]*\r?\n?", re.MULTILINE
)
_INLINE_DATA_URI_IMAGE = re.compile(rb"!\[[^\]\n]{0,200}\]\(<?data:[^\s)>]*>?\)")
_REFERENCED_IMAGE = re.compile(rb"!\[[^\]\n]{0,200}\]\[image\d{1,6}\]")


def markdown_without_inline_images(content: bytes) -> bytes:
    content = _DATA_URI_DEFINITION.sub(b"", content)
    content = _INLINE_DATA_URI_IMAGE.sub(b"[image]", content)
    return _REFERENCED_IMAGE.sub(b"[image]", content)


def parse_live_document(content: bytes, mime_type: str) -> ParsedText:
    """Live lane: ``parse_document`` plus CSV. The selected lane never reads CSV."""
    if mime_type == "text/csv":
        return parse_live_csv(content)
    if mime_type == XLSX:
        return parse_live_xlsx(content)
    if mime_type == DOCX:
        parsed = parse_document(content, mime_type)
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            omitted = any(
                name.casefold().startswith(
                    ("word/media/", "word/header", "word/footer", "word/footnotes", "word/endnotes")
                )
                for name in archive.namelist()
            )
        return ParsedText(parsed.pages, parsed.truncated or omitted)
    if mime_type == "text/markdown" and content and len(content) <= MAX_INPUT_BYTES:
        reduced = markdown_without_inline_images(content)
        parsed = parse_document(reduced, mime_type)
        return ParsedText(
            parsed.pages, parsed.truncated or reduced != content or b"[image]" in reduced
        )
    return parse_document(content, mime_type)


def main() -> None:
    # These untrusted-content leaves fail closed without an OS resource
    # boundary. A parent wall-clock timeout alone is not a memory/CPU sandbox.
    if resource is None:
        sys.stdout.write(json.dumps({"error": "processor_unavailable"}))
        return
    resource.setrlimit(resource.RLIMIT_CORE, (0, 0))
    resource.setrlimit(resource.RLIMIT_CPU, (20, 20))
    if sys.platform == "linux":
        resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024,) * 2)
    try:
        mime = sys.argv[1]
        parser = parse_live_document if mime.startswith("live:") else parse_document
        input_limit = (
            MAX_XLSX_INPUT_BYTES if mime.removeprefix("live:") == XLSX else MAX_INPUT_BYTES
        )
        parsed = parser(sys.stdin.buffer.read(input_limit + 1), mime.removeprefix("live:"))
        result = {"pages": parsed.pages, "truncated": parsed.truncated}
    except ParseError as error:
        result = {"error": str(error)}
    except Exception:
        result = {"error": "invalid_document"}
    sys.stdout.write(json.dumps(result))


if __name__ == "__main__":
    main()
